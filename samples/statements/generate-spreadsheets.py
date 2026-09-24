"""Builds the spreadsheet fixtures for spec 011.

Run from the repository root with openpyxl and xlwt installed:

    python samples/statements/generate-spreadsheets.py

Dates are written as real date cells and amounts as real number cells, the way a
bank's Excel export stores them; `text.xlsx` is the exception, with everything as text.
"""

import datetime
import pathlib
import shutil

import openpyxl
import xlwt

ROOT = pathlib.Path(__file__).resolve().parents[2]
SAMPLES = ROOT / "samples" / "statements"
UNIT = ROOT / "api.tests" / "Fixtures" / "Spreadsheets"
E2E = ROOT / "web" / "e2e" / "fixtures"
FIXED = datetime.datetime(2026, 9, 1)

# Banco do Brasil-like: a title and an account line above the table, a blank line,
# balance lines without a value, and a description split over two columns.
BB_PREAMBLE = [
    ["Banco do Brasil - Extrato de Conta Corrente"],
    ["Agência: 1234-5  Conta: 67890-1  Período: 01/08/2026 a 31/08/2026"],
    [],
]
BB_HEADER = ["Data", "Lançamento", "Detalhes", "Valor (R$)"]
BB_ROWS = [
    [datetime.date(2026, 7, 31), "Saldo Anterior", "", None],
    [datetime.date(2026, 8, 1), "Pix - Recebido", "ACME TECNOLOGIA LTDA", 4500.00],
    [datetime.date(2026, 8, 3), "Compra com Cartão", "SUPERMERCADO ZONA SUL", -187.43],
    [datetime.date(2026, 8, 5), "Pagamento de Boleto", "CONDOMINIO EDIFICIO PRIMAVERA", -680.00],
    [datetime.date(2026, 8, 8), "Pix - Enviado", "JOÃO DA SILVA", -250.00],
    [datetime.date(2026, 8, 12), "Tarifa Pacote de Serviços", "", -39.90],
    [datetime.date(2026, 8, 15), "Compra com Cartão", "POSTO IPIRANGA", -212.35],
    # Stored exactly as Excel stores a sum of 0,10 and 0,20: 0.30000000000000004.
    [datetime.date(2026, 8, 20), "Rende Fácil", "", 0.1 + 0.2],
    [datetime.date(2026, 8, 25), "Pix - Recebido", "MARIA SOUZA", 1234.56],
    [datetime.date(2026, 8, 31), "S A L D O", "", None],
]

TYPED = [
    ["Extrato"],
    ["Conta 123"],
    [],
    ["Data", "Descrição", "Valor"],
    [datetime.date(2026, 8, 3), "Uber", -58],
    [datetime.date(2026, 8, 5), "Mercado", 1234.56],
    [],
    [datetime.date(2026, 8, 10), "Rendimento", 0.1 + 0.2],
]
TEXT = [
    ["Data", "Descrição", "Valor"],
    ["05/08/2026", "Loja", "1.234,56"],
]


def write_xlsx(path, sheets):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.properties.created = FIXED
    workbook.properties.modified = FIXED
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                if value is None or value == "":
                    continue
                cell = sheet.cell(row=r, column=c, value=value)
                if isinstance(value, datetime.date):
                    cell.number_format = "dd/mm/yyyy"
                elif isinstance(value, float):
                    cell.number_format = "#,##0.00"
    workbook.save(path)


def write_xls(path, rows):
    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet("Extrato")
    date_style = xlwt.easyxf(num_format_str="dd/mm/yyyy")
    money_style = xlwt.easyxf(num_format_str="#,##0.00")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value is None or value == "":
                continue
            if isinstance(value, datetime.date):
                sheet.write(r, c, value, date_style)
            elif isinstance(value, float):
                sheet.write(r, c, value, money_style)
            else:
                sheet.write(r, c, value)
    workbook.save(str(path))


def main():
    UNIT.mkdir(parents=True, exist_ok=True)
    bb = BB_PREAMBLE + [BB_HEADER] + BB_ROWS
    write_xlsx(SAMPLES / "bb-extrato-2026-08.xlsx", [("Extrato", bb)])
    write_xls(SAMPLES / "bb-extrato-2026-08.xls", bb)
    shutil.copyfile(SAMPLES / "bb-extrato-2026-08.xlsx", E2E / "bb-extrato.xlsx")
    shutil.copyfile(SAMPLES / "bb-extrato-2026-08.xlsx", UNIT / "bb-extrato.xlsx")

    write_xlsx(UNIT / "typed.xlsx", [("Extrato", TYPED)])
    write_xls(UNIT / "typed.xls", TYPED)
    write_xlsx(UNIT / "text.xlsx", [("Extrato", TEXT)])
    write_xlsx(UNIT / "second-sheet.xlsx", [("Vazia", []), ("Extrato", TYPED[3:6])])
    write_xlsx(UNIT / "empty.xlsx", [("Vazia", [])])


if __name__ == "__main__":
    main()
